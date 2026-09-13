#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Limpia y migra el Excel "Embudo y Registro Leads El Cielo Adventures" a CSV
listos para cargar en la base de datos (ver ../db/schema.sql y ../db/cargar_csv.sql).

Uso:
    python3 limpiar_excel.py "ruta/Embudo_y_Registro_Leads.xlsx" [--salida salida] [--dias-activo 30]

Genera en la carpeta de salida:
    leads.csv, seguimientos.csv, cabanas.csv, reservaciones.csv, gastos_reserva.csv,
    eventos.csv, inscripciones.csv, pasos_embudo.csv, seed_pasos_embudo.sql, reporte.md

La carpeta de salida contiene teléfonos de clientes: está ignorada por git y nunca
debe subirse al repositorio. Requiere: pip install openpyxl
"""
import argparse
import collections
import csv
import datetime as dt
import re
import sys
import unicodedata
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("Falta openpyxl: pip install openpyxl")

HOY = dt.date.today()
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
         "agosto", "septiembre", "octubre", "noviembre", "diciembre"]

# ---------------------------------------------------------------------------
# Embudo nuevo (hoja "Embudo Nuevo"), vuelto datos.
# Cada paso: código, nombre en la hoja (para tomar la plantilla), días de espera
# antes de la siguiente acción, si los días son hábiles, paso siguiente si
# contesta, paso siguiente si no contesta, si es final y cómo cierra el lead.
# "A|B" significa: A si el lead va a Gómez Farías, B si va a San José.
# ---------------------------------------------------------------------------
EMBUDO = [
    # codigo, nombre_hoja,                   espera, habiles, si_contesta, si_no, final, cierra,     orden
    ("1",   "Mensaje Inicial",                1, False, "2|3", "2.1", False, None,        1),
    ("2",   "Gómez Farías",                   1, False, "5",   "2.1", False, None,        2),
    ("3",   "San José",                       1, False, "6",   "2.1", False, None,        3),
    ("4",   "Preguntar",                      1, False, "5|6", "2.1", False, None,        4),
    ("5",   "Cotización Gómez Farías",        1, False, "7",   "3.1", False, None,        5),
    ("6",   "Cotización San José",            1, False, "7",   "3.1", False, None,        6),
    ("7",   "Recomendación",                  1, False, "8",   "4.1", False, None,        7),
    ("8",   "Métodos de pago + Imágen",       1, False, "9",   "4.1", False, None,        8),
    ("9",   "Instrucciones",                  0, False, None,  None,  True,  "reservo",   9),
    ("2.1", "San José ó Gómez Farías",        1, False, "5|6", "3.1", False, None,       10),
    ("3.1", "Contenido de Valor",             2, False, "8",   "4.1", False, None,       11),
    ("4.1", "Leña + Biciletas",               3, False, "8",   "5.1", False, None,       12),
    ("5.1", "Oferta",                         3, False, "8",   "6.1", False, None,       13),
    ("6.1", "Cancelado",                      0, False, None,  None,  True,  "cancelado", 14),
]
# Plantillas de la hoja que no forman parte de la cadena pero conviene conservar.
PLANTILLAS_EXTRA = [
    ("promo-tour", "Tour Gratis"),
    ("promo-10",   "Promo 10% Descuento"),
]

# Cómo se traduce un paso del embudo anterior (el que usa el registro) al nuevo.
# Solo se aplica a leads activos; los cerrados conservan el código original en
# paso_original y quedan sin paso actual.
MAPA_PASO_ANTERIOR = {
    "1":   "1",     # Mensaje inicial
    "1.1": "2.1",   # Llamada -> recontacto
    "2":   "2|3",   # Tipo de viaje / imagen -> info del destino
    "2.1": "2.1",   # Autoridad / infografía -> recontacto
    "3":   "2|3",   # Pregunta del 4x4 / audio -> info del destino
    "3.1": "3.1",   # Urgencia / tour gratis -> contenido de valor
    "4":   "2|3",   # Fotos / reel -> info del destino
    "4.1": "5.1",   # Promoción -> oferta
    "5":   "4",     # Audio -> preguntar antes de cotizar
    "6":   "5|6",   # Cotización enviada
    "7":   "8",     # Forma de pago
}
PASO_ANTERIOR_CANCELADO = "5.1"

# Catálogo de cabañas y cómo se escribieron en el Excel.
CABANAS_BASE = [
    # nombre, ubicacion, capacidad, precio_base
    ("Cabaña San José", "San José", 6, 8299),
    ("Cabaña Alpina Gómez Farías", "Gómez Farías", 4, 2500),
    ("Cabaña Yussef", "San José", None, None),
    ("Cabaña Oliver", "San José", None, None),
    ("Cabaña Cristal", "Gómez Farías", None, None),
    ("Camping", "San José", None, None),
]


# ---------------------------------------------------------------------------
# Utilidades de limpieza
# ---------------------------------------------------------------------------
def quitar_acentos(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def texto(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    return re.sub(r"\s+", " ", str(v)).strip()


def vacio(v):
    return texto(v) == "" or texto(v).upper() in ("X", "-", "N/A", "NA", "/", "XX")


def a_fecha(v):
    if isinstance(v, dt.datetime):
        return v.date() if 2015 <= v.year <= 2035 else None
    if isinstance(v, dt.date):
        return v
    return None


def a_entero(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(v)
    m = re.search(r"\d+", texto(v))
    return int(m.group()) if m else None


def a_dinero(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    t = texto(v).replace(",", "").replace("$", "").replace("MXN", "")
    m = re.search(r"-?\d+(\.\d+)?", t)
    return round(float(m.group()), 2) if m else None


def a_celular(v):
    """Devuelve (lada de país, número de 10 dígitos) o None si no parece un celular."""
    d = re.sub(r"\D", "", texto(v))
    if len(d) == 13 and d.startswith("521"):
        return ("52", d[3:])
    if len(d) == 12 and d.startswith("52"):
        return ("52", d[2:])
    if len(d) == 11 and d.startswith("1"):
        return ("1", d[1:])  # Estados Unidos (leads de Texas)
    if len(d) == 10:
        return ("52", d)
    return None


def a_mes(v):
    if isinstance(v, dt.datetime):
        return MESES[v.month - 1].capitalize()
    t = quitar_acentos(texto(v).lower())
    for m in MESES:
        if t.startswith(quitar_acentos(m)[:3]) and len(t) >= 3:
            return m.capitalize()
    return ""


def a_destino(v):
    """Devuelve (destino normalizado, interés en tours)."""
    t = quitar_acentos(texto(v).upper())
    tours = bool(re.search(r"TOUR|CABALLO|SENDER|\bACT\b|TRANSPORT", t))
    if "JOS" in t:
        return "San José", tours
    if "GOM" in t or "FAR" in t:
        return "Gómez Farías", tours
    if "CAMP" in t:
        return "Camping", tours
    if "EVENTO" in t:
        return "Evento", tours
    if "TOUR" in t or "CIELO" in t:
        return "Tours", tours
    return "Sin definir", tours


def a_paso(v):
    """'5.0' -> '5'; 3.1 -> '3.1'; fecha 2026-01-05 -> '5.1' (Sheets convirtió '5.1' en fecha)."""
    if isinstance(v, dt.datetime):
        return f"{v.day}.{v.month}" if v.year >= 2000 else None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return str(int(v)) if float(v).is_integer() else f"{v:.1f}"
    t = texto(v).replace(",", ".")
    return t if re.fullmatch(r"\d+(\.\d)?", t) else None


def a_bool(v):
    t = quitar_acentos(texto(v).upper())
    if t in ("SI", "S", "PAGADO", "YES"):
        return True
    if t in ("NO", "N"):
        return False
    return None


def a_cabana(v):
    t = quitar_acentos(texto(v).upper())
    if not t or t in ("X", "-", "OTRA"):
        return None
    if "YUSSEF" in t or "YUSEF" in t:
        return "Cabaña Yussef"
    if "OLIVER" in t:
        return "Cabaña Oliver"
    if "CRISTAL" in t:
        return "Cabaña Cristal"
    if "CAMP" in t:
        return "Camping"
    if "JOS" in t:
        return "Cabaña San José"
    if "GOM" in t or "FAR" in t:
        return "Cabaña Alpina Gómez Farías"
    return texto(v).title()


def resolver_destino(codigo, destino):
    """Aplica la regla 'A|B' según el destino del lead."""
    if not codigo or "|" not in codigo:
        return codigo
    a, b = codigo.split("|", 1)
    return b if destino == "San José" else a


def sql_str(v):
    if v is None or v == "":
        return "null"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return str(v)
    return "$q$" + str(v).replace("$q$", "$ q$") + "$q$"


# ---------------------------------------------------------------------------
# Lectura de hojas
# ---------------------------------------------------------------------------
def leer_pasos(wb, rep):
    ws = wb["Embudo Nuevo"]
    plantillas = {}
    for r in range(1, ws.max_row + 1):
        nombre = texto(ws.cell(row=r, column=5).value)
        plantilla = ws.cell(row=r, column=6).value
        if nombre and plantilla is not None:
            plantillas[quitar_acentos(nombre.lower())] = str(plantilla).strip()

    pasos = []
    faltantes = []
    for codigo, nombre_hoja, espera, habiles, si, no, final, cierra, orden in EMBUDO:
        tpl = plantillas.get(quitar_acentos(nombre_hoja.lower()))
        if tpl is None:
            faltantes.append(nombre_hoja)
        pasos.append(dict(codigo=codigo, nombre=nombre_hoja.strip(), dias_espera=espera, dias_habiles=habiles,
                          siguiente_si_contesta=si, siguiente_si_no_contesta=no, es_final=final,
                          cierra_como=cierra, orden=orden, activo=True, plantilla=tpl or ""))
    for i, (codigo, nombre_hoja) in enumerate(PLANTILLAS_EXTRA, start=100):
        tpl = plantillas.get(quitar_acentos(nombre_hoja.lower()))
        pasos.append(dict(codigo=codigo, nombre=nombre_hoja, dias_espera=3, dias_habiles=True,
                          siguiente_si_contesta="8", siguiente_si_no_contesta="6.1", es_final=False,
                          cierra_como=None, orden=i, activo=False, plantilla=tpl or ""))
    rep["pasos"] = dict(total=len(pasos), sin_plantilla=faltantes)
    return pasos


def leer_leads(wb, rep, dias_activo):
    ws = wb["Registro leads"]
    filas = []
    invalidos = []
    conteo_destino_original = collections.Counter()
    conteo_paso_original = collections.Counter()
    pasos_fecha = 0
    for r in range(2, ws.max_row + 1):
        c = lambda col: ws.cell(row=r, column=col).value  # noqa: E731
        if all(vacio(c(col)) for col in range(2, 14)):
            continue
        tel = a_celular(c(3))
        if not tel:
            invalidos.append((r, texto(c(3))))
            continue
        lada, cel = tel
        destino, tours = a_destino(c(4))
        conteo_destino_original[texto(c(4)).upper() or "(vacío)"] += 1
        paso_raw = c(8)
        paso_orig = a_paso(paso_raw)
        if isinstance(paso_raw, dt.datetime) and paso_orig:
            pasos_fecha += 1
        conteo_paso_original[paso_orig or "(sin paso)"] += 1
        reservo = texto(c(13))
        comentarios = texto(c(12))
        com_u = quitar_acentos(comentarios.upper())
        if re.match(r"^S[IÍ]$", quitar_acentos(reservo.upper())) or reservo.startswith("$") or \
                re.search(r"COMPR|RESERVO|RESERVADO|SE COMPLETO", com_u):
            estado = "reservo"
        elif paso_orig == PASO_ANTERIOR_CANCELADO or re.search(r"CANCEL|NO QUIERE", com_u):
            estado = "cancelado"
        else:
            estado = "activo"
        fecha_reg = a_fecha(c(2))
        fecha_paso = a_fecha(c(9)) or fecha_reg
        proxima = a_fecha(c(10))
        if estado == "activo":
            referencia = proxima or fecha_paso
            if referencia is None or referencia < HOY - dt.timedelta(days=dias_activo):
                estado = "archivado"
        filas.append(dict(
            fila=r, lada=lada, celular=cel, destino=destino, interes_tours=tours, mes_viaje=a_mes(c(5)),
            personas=a_entero(c(6)) if not vacio(c(6)) else None,
            noches=a_entero(c(7)) if not vacio(c(7)) else None,
            monto=a_dinero(c(11)) if not vacio(c(11)) else None,
            paso_original=paso_orig, fecha_paso=fecha_paso, proxima=proxima, estado=estado,
            comentarios=comentarios if not vacio(comentarios) else "",
            reservo_txt=reservo if not vacio(reservo) else "",
            creado=fecha_reg or fecha_paso,
        ))

    # Unir duplicados por celular: el registro más reciente manda.
    por_cel = collections.defaultdict(list)
    for f in filas:
        por_cel[(f["lada"], f["celular"])].append(f)
    leads, seguimientos = [], []
    duplicados = 0
    prioridad = {"reservo": 3, "activo": 2, "archivado": 1, "cancelado": 0}
    for (lada, cel), grupo in por_cel.items():
        grupo.sort(key=lambda f: (f["fecha_paso"] or dt.date.min, f["fila"]))
        base = dict(grupo[-1])
        if len(grupo) > 1:
            duplicados += len(grupo) - 1
            base["estado"] = max((g["estado"] for g in grupo), key=lambda e: prioridad[e])
            coms = []
            for g in grupo:
                for parte in (g["comentarios"], g["reservo_txt"]):
                    if parte and parte not in coms:
                        coms.append(parte)
            base["comentarios"] = " / ".join(coms)
            base["creado"] = min((g["creado"] for g in grupo if g["creado"]), default=base["creado"])
            for k in ("destino", "personas", "noches", "monto", "mes_viaje"):
                if base[k] in (None, "", "Sin definir"):
                    for g in reversed(grupo):
                        if g[k] not in (None, "", "Sin definir"):
                            base[k] = g[k]
                            break
        elif base["reservo_txt"] and base["reservo_txt"] not in base["comentarios"]:
            base["comentarios"] = (base["comentarios"] + " / " + base["reservo_txt"]).strip(" /")
        if base["estado"] == "activo":
            paso_actual = resolver_destino(MAPA_PASO_ANTERIOR.get(base["paso_original"] or "", "1"), base["destino"])
        else:
            paso_actual = None
        leads.append(dict(
            lada_pais=lada, celular=cel, nombre="", destino=base["destino"], interes_tours=base["interes_tours"],
            mes_viaje=base["mes_viaje"], personas=base["personas"], noches=base["noches"],
            monto_cotizado=base["monto"], paso_actual=paso_actual, paso_original=base["paso_original"],
            fecha_paso=base["fecha_paso"], proxima_fecha=base["proxima"] if base["estado"] == "activo" else None,
            estado=base["estado"], enviado=True, comentarios=base["comentarios"], creado_en=base["creado"],
            fila_origen="Registro leads fila " + ",".join(str(g["fila"]) for g in grupo),
        ))
        for g in grupo:
            seguimientos.append(dict(
                lada_pais=lada, celular=cel, fecha=g["fecha_paso"] or g["creado"] or HOY,
                paso=resolver_destino(MAPA_PASO_ANTERIOR.get(g["paso_original"] or "", ""), g["destino"]) or None,
                paso_original=g["paso_original"] or "", contesto=None,
                nota=("Migrado del Excel (fila %d, paso %s)%s" % (g["fila"], g["paso_original"] or "sin paso",
                                                                (": " + g["comentarios"]) if g["comentarios"] else "")),
            ))
    rep["leads"] = dict(
        filas=len(filas), invalidos=invalidos, duplicados=duplicados, total=len(leads),
        pasos_convertidos_de_fecha=pasos_fecha, destinos_original=conteo_destino_original,
        destinos=collections.Counter(l["destino"] for l in leads),
        pasos_original=conteo_paso_original, estados=collections.Counter(l["estado"] for l in leads),
        pasos_activos=collections.Counter(l["paso_actual"] for l in leads if l["estado"] == "activo"),
        paso_67_estados=collections.Counter((l["paso_original"], l["estado"]) for l in leads if l["paso_original"] in ("6", "7")),
    )
    return leads, seguimientos


def leer_reservaciones(wb, rep):
    ws = wb["Reservaciones "]
    reservas, gastos = [], []
    sin_telefono = 0
    fechas_corregidas = []
    conceptos = {24: "Gastos operativos", 25: "Sueldo personal", 26: "Mantenimiento 250",
                 27: "Mantenimiento 300", 28: "Marketing", 29: "Comisión 10%"}
    for r in range(2, ws.max_row + 1):
        c = lambda col: ws.cell(row=r, column=col).value  # noqa: E731
        nombre = texto(c(3))
        if vacio(nombre) or nombre.upper() == "NOMBRE":
            continue
        tel = a_celular(c(4))
        if not tel:
            sin_telefono += 1
        lada, cel = tel if tel else ("52", None)
        llegada, salida = a_fecha(c(8)), a_fecha(c(9))
        nota_fechas = ""
        if llegada and salida and salida < llegada:
            # "31 al 1": la salida se capturó con el mes equivocado; se prueba el mes siguiente.
            anio, mes_sig = (salida.year + 1, 1) if salida.month == 12 else (salida.year, salida.month + 1)
            try:
                candidata = salida.replace(year=anio, month=mes_sig)
            except ValueError:
                candidata = None
            if candidata and candidata > llegada:
                nota_fechas = "Fecha de salida corregida (venía %s)" % salida.isoformat()
                salida = candidata
            else:
                nota_fechas = "Fecha de salida inválida en el Excel (%s), se dejó vacía" % salida.isoformat()
                salida = None
            fechas_corregidas.append((r, nota_fechas))
        total = a_dinero(c(12)) if not vacio(c(12)) else (a_dinero(c(22)) if not vacio(c(22)) else None)
        anticipo = a_dinero(c(13)) if not vacio(c(13)) else None
        liquido = a_fecha(c(15))
        liquidada = liquido is not None or quitar_acentos(texto(c(15)).upper()) == "PAGADO" or \
            (total is not None and anticipo is not None and anticipo >= total)
        if salida and salida < HOY:
            estado = "completada"
        elif liquidada:
            estado = "liquidada"
        else:
            estado = "apartada"
        coms = [texto(c(col)) for col in (21, 33) if not vacio(c(col)) and texto(c(col)).upper() != "COMENTARIOS"]
        if nota_fechas:
            coms.append(nota_fechas)
        ubic, _ = a_destino(c(7))
        reservas.append(dict(
            fila_origen="Reservaciones fila %d" % r, nombre=nombre.title(), lada_pais=lada, celular=cel,
            liquidada=liquidada,
            personas=a_entero(c(5)) if not vacio(c(5)) else None, cabana=a_cabana(c(6)), ubicacion=ubic,
            fecha_llegada=llegada, fecha_salida=salida, fechas_texto=texto(c(10)) if not vacio(c(10)) else "",
            mes=a_mes(c(11)), total=total, anticipo=anticipo, fecha_anticipo=a_fecha(c(14)),
            fecha_liquidacion=liquido, hora_llegada=texto(c(16)) if not vacio(c(16)) else "",
            ocupa_transporte=a_bool(c(17)), estado_transporte=texto(c(18)) if not vacio(c(18)) else "",
            limpieza_rapida=a_bool(c(19)), dias_limpiar=a_entero(c(20)) if not vacio(c(20)) else None,
            estado=estado, comentarios=" / ".join(coms),
        ))
        for col, concepto in conceptos.items():
            monto = a_dinero(c(col)) if not vacio(c(col)) else None
            if monto:
                gastos.append(dict(fila_origen="Reservaciones fila %d" % r, concepto=concepto, monto=monto))
    rep["reservaciones"] = dict(total=len(reservas), sin_telefono=sin_telefono, gastos=len(gastos),
                                sin_fecha=sum(1 for x in reservas if not x["fecha_llegada"]),
                                fechas_corregidas=fechas_corregidas,
                                estados=collections.Counter(x["estado"] for x in reservas),
                                cabanas=collections.Counter(x["cabana"] or "(sin cabaña)" for x in reservas))
    return reservas, gastos


def leer_eventos(wb, rep):
    ws = wb["Eventos"]
    etiquetas = {}
    for r in range(1, ws.max_row + 1):
        etiquetas[quitar_acentos(texto(ws.cell(row=r, column=2).value).lower())] = r
    fila = lambda clave: etiquetas.get(clave)  # noqa: E731
    campos = dict(nombre=fila("nombre del evento"), mes=fila("mes"), fecha=fila("fecha de evento"),
                  lugar=fila("lugar del evento"), fecha_pub=fila("fecha de publicacion"),
                  incluye=fila("que incluye el viaje"), costo=fila("costo"), precio=fila("precio al publico"),
                  respuesta=fila("respuesta predeterminada"),
                  extra=fila("como hacer de este viaje algo extraordinario"))
    eventos = []
    vistos = set()
    for col in range(3, ws.max_column + 1):
        v = lambda k: ws.cell(row=campos[k], column=col).value if campos.get(k) else None  # noqa: E731
        nombre = texto(v("nombre"))
        if vacio(nombre) or nombre in vistos:
            continue
        vistos.add(nombre)
        lugar, _ = a_destino(v("lugar"))
        eventos.append(dict(
            nombre=nombre, mes=a_mes(v("mes")), fecha=a_fecha(v("fecha")), lugar=lugar,
            fecha_publicacion=a_fecha(v("fecha_pub")), incluye=texto(v("incluye")),
            costo=a_dinero(v("costo")), precio_publico=a_dinero(v("precio")),
            respuesta_predeterminada=texto(v("respuesta")), extra=texto(v("extra")),
        ))
    # Eventos con inscritos que no aparecen en la programación
    ws2 = wb["Reservaciones Eventos"]
    inscripciones = []
    for r in range(5, ws2.max_row + 1):
        c = lambda col: ws2.cell(row=r, column=col).value  # noqa: E731
        nombre, evento = texto(c(3)), texto(c(5))
        if vacio(nombre) or vacio(evento) or nombre.upper() == "NOMBRE" or a_dinero(evento) is not None:
            continue
        tel = a_celular(c(4))
        lada, cel = tel if tel else ("52", None)
        pagado, pendiente = a_dinero(c(7)) or 0, a_dinero(c(8)) or 0
        inscripciones.append(dict(
            fila_origen="Reservaciones Eventos fila %d" % r,
            evento=evento, nombre=nombre.title(), lada_pais=lada, celular=cel,
            precio=pagado + pendiente, monto_pagado=pagado, fecha_inscripcion=a_fecha(c(9)),
            alojamiento=texto(c(10)) if not vacio(c(10)) else "", paga_alojamiento=a_bool(c(11)),
            monto_alojamiento=a_dinero(c(12)) if not vacio(c(12)) else None,
            cuenta=texto(c(15)) if not vacio(c(15)) else "", comentarios=texto(c(14)) if not vacio(c(14)) else "",
        ))
    nombres = {e["nombre"] for e in eventos}
    for ins in inscripciones:
        if ins["evento"] not in nombres:
            eventos.append(dict(nombre=ins["evento"], mes="", fecha=None, lugar="San José", fecha_publicacion=None,
                                incluye="", costo=None, precio_publico=None, respuesta_predeterminada="", extra=""))
            nombres.add(ins["evento"])
    rep["eventos"] = dict(total=len(eventos), inscripciones=len(inscripciones),
                          por_evento=collections.Counter(i["evento"] for i in inscripciones))
    return eventos, inscripciones


# ---------------------------------------------------------------------------
# Escritura
# ---------------------------------------------------------------------------
def escribir_csv(ruta, filas, columnas):
    with open(ruta, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=columnas, extrasaction="ignore")
        w.writeheader()
        for fila in filas:
            w.writerow({k: ("" if fila.get(k) is None else fila.get(k)) for k in columnas})


def escribir_seed(ruta, pasos):
    lineas = ["-- Generado por limpiar_excel.py a partir de la hoja 'Embudo Nuevo'. Editable.",
              "insert into pasos_embudo (codigo, nombre, plantilla, dias_espera, dias_habiles, siguiente_si_contesta,",
              "  siguiente_si_no_contesta, es_final, cierra_como, orden, activo) values"]
    valores = []
    for p in pasos:
        valores.append("  (%s)" % ", ".join(sql_str(p[k]) for k in (
            "codigo", "nombre", "plantilla", "dias_espera", "dias_habiles", "siguiente_si_contesta",
            "siguiente_si_no_contesta", "es_final", "cierra_como", "orden", "activo")))
    lineas.append(",\n".join(valores))
    lineas.append("on conflict (codigo) do update set nombre = excluded.nombre, plantilla = excluded.plantilla,")
    lineas.append("  dias_espera = excluded.dias_espera, dias_habiles = excluded.dias_habiles,")
    lineas.append("  siguiente_si_contesta = excluded.siguiente_si_contesta,")
    lineas.append("  siguiente_si_no_contesta = excluded.siguiente_si_no_contesta, es_final = excluded.es_final,")
    lineas.append("  cierra_como = excluded.cierra_como, orden = excluded.orden, activo = excluded.activo;")
    Path(ruta).write_text("\n".join(lineas) + "\n", encoding="utf-8")


def tabla(conteo, titulo_col="Valor", limite=None):
    items = conteo.most_common(limite) if hasattr(conteo, "most_common") else list(conteo.items())
    out = ["| %s | Registros |" % titulo_col, "|---|---:|"]
    out += ["| %s | %d |" % (k, v) for k, v in items]
    return "\n".join(out)


def escribir_reporte(ruta, rep, dias_activo, origen):
    L, R, E, P = rep["leads"], rep["reservaciones"], rep["eventos"], rep["pasos"]
    partes = [
        "# Reporte de migración",
        "",
        "Archivo: `%s`  " % origen,
        "Fecha: %s  " % HOY.isoformat(),
        "Un lead se considera activo si su próxima fecha (o su último cambio) es de los últimos %d días." % dias_activo,
        "",
        "## Leads",
        "",
        "| Concepto | Cantidad |", "|---|---:|",
        "| Filas leídas en Registro leads | %d |" % L["filas"],
        "| Filas descartadas por celular inválido | %d |" % len(L["invalidos"]),
        "| Filas unidas por celular repetido | %d |" % L["duplicados"],
        "| Leads resultantes | %d |" % L["total"],
        "| Pasos que Sheets había convertido en fecha y se recuperaron | %d |" % L["pasos_convertidos_de_fecha"],
        "",
        "### Estado de los leads", "", tabla(L["estados"], "Estado"),
        "",
        "### Paso actual de los leads activos (ya en el embudo nuevo)", "", tabla(L["pasos_activos"], "Paso"),
        "",
        "### Destino normalizado", "", tabla(L["destinos"], "Destino"),
        "",
        "### Cómo estaba escrito el destino (%d variantes)" % len(L["destinos_original"]), "",
        tabla(L["destinos_original"], "Texto original", 40),
        "",
        "### Paso del embudo anterior tal como venía", "", tabla(L["pasos_original"], "Paso"),
        "",
        "### Celulares inválidos (fila, texto)", "",
        "\n".join("- fila %d: `%s`" % (r, t) for r, t in L["invalidos"]) or "- ninguno",
        "",
        "## Reservaciones",
        "",
        "| Concepto | Cantidad |", "|---|---:|",
        "| Reservaciones leídas | %d |" % R["total"],
        "| Sin teléfono válido | %d |" % R["sin_telefono"],
        "| Sin fecha de llegada | %d |" % R["sin_fecha"],
        "| Gastos capturados | %d |" % R["gastos"],
        "", tabla(R["estados"], "Estado"), "", tabla(R["cabanas"], "Cabaña"),
        "",
        "### Fechas corregidas", "",
        "\n".join("- fila %d: %s" % (r, n) for r, n in R["fechas_corregidas"]) or "- ninguna",
        "",
        "## Eventos",
        "",
        "| Concepto | Cantidad |", "|---|---:|",
        "| Eventos | %d |" % E["total"],
        "| Inscripciones | %d |" % E["inscripciones"],
        "", tabla(E["por_evento"], "Evento"),
        "",
        "## Embudo",
        "",
        "Pasos generados: %d. Sin plantilla en la hoja: %s" % (P["total"], ", ".join(P["sin_plantilla"]) or "ninguno"),
        "",
        "## Supuestos que conviene confirmar",
        "",
        "- El paso `5.1` del registro se tomó como *cancelado* (así estaba definido en los embudos anteriores).",
        "- Los pasos `6` y `7` del registro se interpretaron como *cotización enviada* y *forma de pago enviada*,",
        "  siguiendo el embudo de San José. Distribución encontrada: %s." % ", ".join(
            "%s/%s: %d" % (p, e, n) for (p, e), n in sorted(L["paso_67_estados"].items())),
        "- Los leads cerrados (reservó, cancelado, archivado) conservan su paso original pero quedan sin paso",
        "  actual; si se reactivan, se les asigna paso desde la aplicación.",
        "- Un lead se marcó como *reservó* si la columna RESERVO? decía Sí o traía un monto, o si el comentario",
        "  decía compró, reservó o se completó la reservación.",
    ]
    Path(ruta).write_text("\n".join(partes) + "\n", encoding="utf-8")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("archivo", help="Excel exportado de Google Sheets (.xlsx)")
    ap.add_argument("--salida", default="salida", help="carpeta de salida (default: salida)")
    ap.add_argument("--dias-activo", type=int, default=30, help="días hacia atrás para considerar un lead activo")
    args = ap.parse_args()

    salida = Path(args.salida)
    salida.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(args.archivo, data_only=True)
    rep = {}

    pasos = leer_pasos(wb, rep)
    leads, seguimientos = leer_leads(wb, rep, args.dias_activo)
    reservas, gastos = leer_reservaciones(wb, rep)
    eventos, inscripciones = leer_eventos(wb, rep)

    cabanas = [dict(nombre=n, ubicacion=u, capacidad=c, precio_base=p) for n, u, c, p in CABANAS_BASE]
    conocidas = {c["nombre"] for c in cabanas}
    for r in reservas:
        if r["cabana"] and r["cabana"] not in conocidas:
            cabanas.append(dict(nombre=r["cabana"], ubicacion=r["ubicacion"], capacidad=None, precio_base=None))
            conocidas.add(r["cabana"])

    escribir_csv(salida / "pasos_embudo.csv", pasos, [
        "codigo", "nombre", "dias_espera", "dias_habiles", "siguiente_si_contesta", "siguiente_si_no_contesta",
        "es_final", "cierra_como", "orden", "activo", "plantilla"])
    escribir_seed(salida / "seed_pasos_embudo.sql", pasos)
    escribir_csv(salida / "leads.csv", leads, [
        "lada_pais", "celular", "nombre", "destino", "interes_tours", "mes_viaje", "personas", "noches", "monto_cotizado",
        "paso_actual", "paso_original", "fecha_paso", "proxima_fecha", "estado", "enviado", "comentarios",
        "creado_en", "fila_origen"])
    escribir_csv(salida / "seguimientos.csv", seguimientos,
                 ["lada_pais", "celular", "fecha", "paso", "paso_original", "contesto", "nota"])
    escribir_csv(salida / "cabanas.csv", cabanas, ["nombre", "ubicacion", "capacidad", "precio_base"])
    escribir_csv(salida / "reservaciones.csv", reservas, [
        "fila_origen", "nombre", "lada_pais", "celular", "personas", "cabana", "ubicacion", "fecha_llegada",
        "fecha_salida", "fechas_texto", "mes", "total", "anticipo", "fecha_anticipo", "liquidada",
        "fecha_liquidacion", "hora_llegada",
        "ocupa_transporte", "estado_transporte", "limpieza_rapida", "dias_limpiar", "estado", "comentarios"])
    escribir_csv(salida / "gastos_reserva.csv", gastos, ["fila_origen", "concepto", "monto"])
    escribir_csv(salida / "eventos.csv", eventos, [
        "nombre", "mes", "fecha", "lugar", "fecha_publicacion", "incluye", "costo", "precio_publico",
        "respuesta_predeterminada", "extra"])
    escribir_csv(salida / "inscripciones.csv", inscripciones, [
        "fila_origen", "evento", "nombre", "lada_pais", "celular", "precio", "monto_pagado", "fecha_inscripcion",
        "alojamiento", "paga_alojamiento", "monto_alojamiento", "cuenta", "comentarios"])
    escribir_reporte(salida / "reporte.md", rep, args.dias_activo, Path(args.archivo).name)

    L = rep["leads"]
    print("Leads: %d filas -> %d leads (%d duplicados unidos, %d celulares inválidos)" % (
        L["filas"], L["total"], L["duplicados"], len(L["invalidos"])))
    print("Estados:", dict(L["estados"]))
    print("Reservaciones: %d  Gastos: %d  Eventos: %d  Inscripciones: %d" % (
        rep["reservaciones"]["total"], rep["reservaciones"]["gastos"], rep["eventos"]["total"], rep["eventos"]["inscripciones"]))
    print("Salida en", salida.resolve())


if __name__ == "__main__":
    main()
